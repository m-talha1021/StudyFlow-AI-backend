from io import BytesIO

from pypdf import PdfReader
from docx import Document
from pptx import Presentation


# Supported file types
ALLOWED_EXTENSIONS = {
    "pdf",
    "docx",
    "pptx"
}


# ------------------------------------------------
# Get File Extension
# ------------------------------------------------

def get_file_extension(filename):

    return filename.rsplit(
        ".",
        1
    )[-1].lower()


# ------------------------------------------------
# Extract PDF Text
# ------------------------------------------------

def extract_pdf(file_bytes):

    reader = PdfReader(
        BytesIO(file_bytes)
    )

    pages = []

    for page in reader.pages:

        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n\n".join(pages)


# ------------------------------------------------
# Extract DOCX Text
# ------------------------------------------------

def extract_docx(file_bytes):

    document = Document(
        BytesIO(file_bytes)
    )

    paragraphs = []

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            paragraphs.append(
                paragraph.text
            )

    return "\n".join(paragraphs)


# ------------------------------------------------
# Extract PPTX Text
# ------------------------------------------------

def extract_pptx(file_bytes):

    presentation = Presentation(
        BytesIO(file_bytes)
    )

    slides = []

    for slide in presentation.slides:

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                if shape.text.strip():

                    slide_text.append(
                        shape.text
                    )

        if slide_text:

            slides.append(
                "\n".join(slide_text)
            )

    return "\n\n".join(slides)


# ------------------------------------------------
# Main Text Extraction Function
# ------------------------------------------------

def extract_text(
    file_bytes,
    filename
):

    extension = get_file_extension(
        filename
    )

    if extension not in ALLOWED_EXTENSIONS:

        raise ValueError(
            "Unsupported file type. "
            "Please upload PDF, DOCX, "
            "or PPTX."
        )


    if extension == "pdf":

        return extract_pdf(
            file_bytes
        )


    if extension == "docx":

        return extract_docx(
            file_bytes
        )


    if extension == "pptx":

        return extract_pptx(
            file_bytes
        )


    raise ValueError(
        "Unsupported file type."
    )